from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent


def hyperlink(paragraph, label, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '185B4B')
    props.append(color)
    under = OxmlElement('w:u')
    under.set(qn('w:val'), 'single')
    props.append(under)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, value):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*.*?\*\*)'
    for token in re.split(pattern, value):
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('[') and '](' in token:
            match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
            hyperlink(paragraph, match[1], match[2])
        else:
            paragraph.add_run(token)


def convert(path):
    compact = path.name.startswith('01-')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.header_distance = sec.footer_distance = Cm(0.7)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10 if compact else 10.5)
    normal.paragraph_format.space_after = Pt(5 if compact else 6)
    normal.paragraph_format.line_spacing = 1.03 if compact else 1.08
    for name, size in [('Title', 18), ('Heading 1', 14), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('185B4B')
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
    lines = path.read_text(encoding='utf-8').splitlines()
    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line)
            r.font.name = 'Consolas'
            r.font.size = Pt(7.5)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [s.strip() for s in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', s) for s in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Table Grid'
            for idx, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    p = cell.paragraphs[0]
                    inline(p, value)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if idx == 0:
                            run.bold = True
                    if idx == 0:
                        shade = OxmlElement('w:shd')
                        shade.set(qn('w:fill'), 'E7EFEA')
                        cell._tc.get_or_add_tcPr().append(shade)
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement('w:cantSplit'))
                if idx == 0:
                    props.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        elif line.startswith('# '):
            inline(doc.add_paragraph(style='Title'), line[2:])
        elif line.startswith('## '):
            inline(doc.add_paragraph(style='Heading 1'), line[3:])
        elif line.startswith('### '):
            inline(doc.add_paragraph(style='Heading 2'), line[4:])
        elif line.startswith('- '):
            inline(doc.add_paragraph(style='List Bullet'), line[2:])
        elif line.strip():
            inline(doc.add_paragraph(), line)
        i += 1
    footer = sec.footer.paragraphs[0]
    footer.add_run('TRÀ VINH DIGITAL HERITAGE · 19/09/2026').font.size = Pt(8)
    doc.core_properties.title = lines[0].lstrip('# ')
    doc.core_properties.subject = 'Tài liệu nền cho prototype bàn tương tác'
    doc.core_properties.language = 'vi-VN'
    target = path.with_suffix('.docx')
    doc.save(target)
    print(target.name)


if __name__ == '__main__':
    for source in sorted(ROOT.glob('0[123]-*.md')):
        convert(source)
